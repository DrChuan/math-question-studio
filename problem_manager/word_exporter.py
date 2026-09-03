from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from problem_manager.latex_support import latex_to_omml, split_math_segments
from problem_manager.models import CHOICE_TYPES, OPTION_LABELS, Question

LATIN_FONT = "Times New Roman"
CJK_FONT = "宋体"
TITLE_CJK_FONT = "黑体"


@dataclass
class WordExportOptions:
    title: str = "高中数学题目"
    include_analysis: bool = False
    include_note: bool = False
    prefer_mathtype: bool = False


def export_questions_to_word(
    questions: list[Question],
    file_path: str | Path,
    options: WordExportOptions | None = None,
) -> Path:
    options = options or WordExportOptions()
    document = Document()
    _configure_document_styles(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _append_plain_run(title, options.title, east_asia_font=TITLE_CJK_FONT, size=Pt(16), bold=True)

    for index, question in enumerate(questions, start=1):
        stem = document.add_paragraph()
        _configure_question_paragraph(stem, index)
        _append_plain_run(stem, f"{index}．")
        if question.question_type == "多选题":
            _append_plain_run(stem, "（多选）")
        _append_mixed_content(stem, question.stem, emit_latex_source=options.prefer_mathtype)

        if question.question_type in CHOICE_TYPES:
            _add_choice_block(document, question, emit_latex_source=options.prefer_mathtype)

        _add_labeled_block(document, "答案", question.answer, emit_latex_source=options.prefer_mathtype)
        if options.include_analysis:
            _add_labeled_block(document, "解析", question.analysis, emit_latex_source=options.prefer_mathtype)
        if options.include_note:
            _add_labeled_block(document, "备注", question.note, emit_latex_source=options.prefer_mathtype)

    output = Path(file_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def _configure_document_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    document.styles["Normal"].paragraph_format.space_after = Pt(0)
    document.styles["Normal"].paragraph_format.line_spacing = 1.25


def _add_labeled_block(document: Document, label: str, content: str, emit_latex_source: bool) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.74)
    paragraph.paragraph_format.first_line_indent = Cm(-0.74)
    _append_plain_run(paragraph, f"{label}：", bold=True)
    _append_mixed_content(paragraph, content, emit_latex_source=emit_latex_source)


def _append_mixed_content(paragraph, text: str, emit_latex_source: bool = False) -> None:
    current_paragraph = paragraph
    for kind, content, _display in split_math_segments(text):
        if kind == "text":
            parts = content.split("\n")
            for index, part in enumerate(parts):
                if index > 0:
                    current_paragraph = paragraph._parent.add_paragraph()
                    current_paragraph.paragraph_format.left_indent = paragraph.paragraph_format.left_indent
                    current_paragraph.paragraph_format.first_line_indent = paragraph.paragraph_format.first_line_indent
                if part:
                    _append_plain_run(current_paragraph, part)
        elif content:
            if emit_latex_source:
                _append_plain_run(current_paragraph, _inline_latex_source(content))
            else:
                current_paragraph._element.append(latex_to_omml(content))


def _inline_latex_source(content: str) -> str:
    return "$" + content.strip() + "$"


def _append_plain_run(
    paragraph,
    text: str,
    latin_font: str = LATIN_FONT,
    east_asia_font: str = CJK_FONT,
    size: Pt | None = None,
    bold: bool = False,
) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    run.font.name = latin_font
    run.font.bold = bold
    run.font.size = size or Pt(12)
    run._element.rPr.rFonts.set(qn("w:ascii"), latin_font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin_font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia_font)


def _configure_question_paragraph(paragraph, index: int) -> None:
    indent = Cm(0.82 + max(0, len(str(index)) - 1) * 0.25)
    paragraph.paragraph_format.left_indent = indent
    paragraph.paragraph_format.first_line_indent = -indent


def _add_choice_block(document: Document, question: Question, emit_latex_source: bool) -> None:
    active_options = [
        (label, _normalize_inline_text(option))
        for label, option in zip(OPTION_LABELS, question.options)
        if option.strip()
    ]
    if not active_options:
        return

    row_size = _decide_choice_row_size([content for _, content in active_options])
    rows = [active_options[i : i + row_size] for i in range(0, len(active_options), row_size)]

    for row in rows:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.9)
        _set_tab_stops(paragraph, len(row), row_size)
        for cell_index, (label, content) in enumerate(row):
            if cell_index > 0:
                _append_tab(paragraph)
            _append_plain_run(paragraph, f"{label}. ")
            _append_mixed_content(paragraph, content, emit_latex_source=emit_latex_source)


def _normalize_inline_text(text: str) -> str:
    compact = text.replace("\n", " ")
    compact = re.sub(r"\s+", " ", compact).strip()
    return compact


def _estimate_text_width(text: str) -> float:
    simplified = re.sub(r"\\[a-zA-Z]+", "x", text)
    simplified = simplified.replace("{", "").replace("}", "").replace("$", "")
    width = 0.0
    for char in simplified:
        width += 1.0 if ord(char) > 127 else 0.55
    return width


def _decide_choice_row_size(options: list[str]) -> int:
    if not options:
        return 1
    widths = [_estimate_text_width(item) for item in options]
    one_row_width = sum(widths) + max(0, len(widths) - 1) * 3.2
    if len(options) <= 4 and max(widths) <= 4.8 and one_row_width <= 22:
        return len(options)

    pair_widths = []
    for start in range(0, len(widths), 2):
        pair_widths.append(sum(widths[start : start + 2]) + (3.2 if start + 1 < len(widths) else 0))
    if max(pair_widths) <= 17:
        return 2

    return 1


def _set_tab_stops(paragraph, visible_count: int, row_size: int) -> None:
    tab_stops = paragraph.paragraph_format.tab_stops
    if row_size >= 4:
        positions = [Cm(3.8), Cm(7.4), Cm(11.0)]
    elif row_size == 3:
        positions = [Cm(4.9), Cm(9.8)]
    elif row_size == 2 and visible_count > 1:
        positions = [Cm(7.4)]
    else:
        positions = []
    for position in positions:
        tab_stops.add_tab_stop(position)


def _append_tab(paragraph) -> None:
    run = paragraph.add_run()
    run.add_tab()
